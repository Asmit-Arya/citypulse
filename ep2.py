import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_full_page_assignment():
    doc = Document()

    # --- Page Setup: A4 with Equal 0.6" Margins on All 4 Sides ---
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Base typography (Pure Monochrome)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0, 0, 0)

    # Helper: Formatted Paragraph with proportional spacing
    def add_p(text="", bold_prefix="", space_before=4, space_after=6, line_spacing=1.2, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.size = Pt(10.5)
            rb.font.color.rgb = RGBColor(0, 0, 0)
        if text:
            rt = p.add_run(text)
            rt.italic = italic
            rt.font.size = Pt(10.5)
            rt.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # Helper: Section Titles
    def add_sec_title(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # Helper: Cell content formatting with internal padding
    def set_cell(cell, bold_label="", body_text="", font_size=9.5, space_before=3, space_after=3, line_spacing=1.18):
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if bold_label:
            rb = p.add_run(bold_label)
            rb.bold = True
            rb.font.size = Pt(font_size)
            rb.font.color.rgb = RGBColor(0, 0, 0)
        if body_text:
            rt = p.add_run(body_text)
            rt.font.size = Pt(font_size)
            rt.font.color.rgb = RGBColor(0, 0, 0)

    # ==========================================================
    # PAGE 1: YOUR IDEAL CUSTOMERS (SEGMENTATION)
    # ==========================================================
    # Top Metadata Box
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.style = 'Table Grid'
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    set_cell(meta_table.rows[0].cells[0], "Course: ", "IT Workshop — Entrepreneurship Module", font_size=10, space_before=3, space_after=3)
    set_cell(meta_table.rows[0].cells[1], "Submission Date: ", "15th Oct 2026", font_size=10, space_before=3, space_after=3)
    set_cell(meta_table.rows[1].cells[0], "Instructor: ", "Dr. Amit Agrawal", font_size=10, space_before=3, space_after=3)
    set_cell(meta_table.rows[1].cells[1], "Startup / Project Name: ", "CityPulse", font_size=10, space_before=3, space_after=3)
    
    merged_cell = meta_table.rows[2].cells[0]
    merged_cell.merge(meta_table.rows[2].cells[1])
    set_cell(merged_cell, "Team Members: ", 
             "1. Asmit Arya (Roll: [Roll No], CSE) | 2. [Member 2 Name] (Roll: [Roll No], CSE) | "
             "3. [Member 3 Name] (Roll: [Roll No], CSE) | 4. [Member 4 Name] (Roll: [Roll No], CSE)",
             font_size=9.5, space_before=4, space_after=4)

    add_sec_title("PAGE 1: YOUR IDEAL CUSTOMERS (CUSTOMER DISCOVERY)")

    add_p("College students living in hostels bound by lecture hours, mandatory labs, and strict campus curfews; "
          "hostellers and local residents who rely on neighborhood shops outside campus (clinics, barbers, and phone repair counters) for everyday needs; "
          "digital natives who use on-demand apps daily but still face blind offline waiting rooms; "
          "busy individuals with narrow 1-hour windows during the day to run essential personal errands.",
          bold_prefix="1. What groups of people do you, or did you, belong to?\n",
          space_before=4, space_after=7, line_spacing=1.22)

    add_p("Small, neighborhood business owners (local doctors, barbers, repair specialists) who lose walk-in customers and revenue whenever "
          "their shops get crowded, alongside students and working folks who cannot afford to waste 1 to 2 hours sitting in stagnant waiting rooms.",
          bold_prefix="2. Who are you passionate about helping?\n",
          space_before=4, space_after=7, line_spacing=1.22)

    add_p("• Potential Segment #1: Local Clinic & OPD Doctors (Small 1-to-2 doctor setups where waiting lines easily pile up)\n"
          "• Potential Segment #2: College Students & Busy Locals (People trying to run errands on tight schedules)\n"
          "• Potential Segment #3: Neighborhood Salons & Repair Shops (Everyday walk-in counters hit by unpredictable evening crowds)",
          bold_prefix="3. Pick the 3 groups you're most excited to help:\n",
          space_before=4, space_after=8, line_spacing=1.22)

    # 3-Segment Comparison Grid
    p1_tbl = doc.add_table(rows=3, cols=3)
    p1_tbl.style = 'Table Grid'
    p1_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_p1 = [
        "Potential Segment #1:\nLocal Clinic & OPD Doctors",
        "Potential Segment #2:\nStudents & Busy Locals",
        "Potential Segment #3:\nSalons & Repair Shops"
    ]
    for i, h in enumerate(headers_p1):
        set_cell(p1_tbl.rows[0].cells[i], h, "", font_size=9.5, space_before=4, space_after=4)
        p1_tbl.rows[0].cells[i].paragraphs[0].runs[0].bold = True

    p1_problems = [
        "• Severe waiting-room congestion during evening peak OPD hours.\n• Patients leaving without consultation ('walk-aways') due to crowded lobbies.\n• Reception staff constantly interrupted by 'how much longer?' inquiries.",
        "• Having zero visibility into live line lengths before traveling to the shop.\n• Wasting 45 to 90 minutes sitting idle on waiting benches.\n• Missed study blocks or delayed schedules from unpredictable service delays.",
        "• Severe demand imbalance: empty chairs during the day, chaotic rushes in the evening.\n• Direct revenue lost when prospective walk-ins see full couches and walk away.\n• Commercial booking apps are too rigid and expensive for casual walk-ins."
    ]
    for i, prob in enumerate(p1_problems):
        set_cell(p1_tbl.rows[1].cells[i], f"{4+i}. Biggest Problems:\n", prob, font_size=9, space_before=4, space_after=4, line_spacing=1.2)

    p1_hardest = [
        "The 6:00 PM – 8:30 PM OPD rush. Over a dozen patients crowd a tiny reception area, disputes emerge over turn sequences, and consultations get rushed.",
        "Attempting to finish a 20-minute errand during an open 1-hour class window, only to arrive and find an unexpected 90-minute queue ahead.",
        "Watching high-paying customers peek inside the front door, see all seats full, turn around, and walk straight to a competitor."
    ]
    for i, hard in enumerate(p1_hardest):
        set_cell(p1_tbl.rows[2].cells[i], "Hardest Part of their Day:\n", hard, font_size=9, space_before=4, space_after=4, line_spacing=1.2)

    doc.add_page_break()

    # ==========================================================
    # PAGE 2: CUSTOMER EMOTIONS & SEGMENT VALIDATION
    # ==========================================================
    add_sec_title("PAGE 2: CUSTOMER EMOTIONS & VALIDATION")

    p2_tbl = doc.add_table(rows=5, cols=4)
    p2_tbl.style = 'Table Grid'
    p2_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers_p2 = [
        "7. Customer Emotions",
        "Segment #1: Clinic Doctors",
        "Segment #2: Students & Patrons",
        "Segment #3: Salons & Repair"
    ]
    for i, h in enumerate(headers_p2):
        set_cell(p2_tbl.rows[0].cells[i], h, "", font_size=9.5, space_before=4, space_after=4)
        p2_tbl.rows[0].cells[i].paragraphs[0].runs[0].bold = True

    emotions_data = [
        ("What is this segment afraid of?",
         "Patients getting hostile over long delays, leaving negative reviews, or switching to corporate hospital OPDs.",
         "Missing assignment deadlines, exam study blocks, or scheduled classes due to being stuck in an unexpected queue.",
         "Regular, loyal customers giving up on the wait and permanently switching to commercial salon chains in malls."),
        ("What frustrates this segment?",
         "Being interrupted repeatedly about wait times while trying to focus on diagnosing and treating patients.",
         "Showing up in person only to hear 'just 10 minutes,' and still sitting on the bench 50 minutes later.",
         "Having chairs sit empty for hours, followed by turning paying customers away during the chaotic evening rush."),
        ("What overwhelms this segment?",
         "Managing messy paper registers, noisy waiting rooms, and impatient patients all at the same time.",
         "Trying to balance intense coursework when every single chore outside campus turns into an unpredictable time sink.",
         "Complicated enterprise booking tools that charge high monthly fees and require bulky desktop setups."),
        ("What does this segment love?",
         "A calm waiting room where patients arrive right on time for their turn, allowing focus on quality care.",
         "Knowing the line beforehand so they can leave their room at the right time and get served right away.",
         "A predictable flow of customers throughout the day, steady chair utilization, and happy recurring clients.")
    ]

    for row_idx, (dimension, s1, s2, s3) in enumerate(emotions_data, start=1):
        set_cell(p2_tbl.rows[row_idx].cells[0], dimension, "", font_size=9, space_before=4, space_after=4)
        p2_tbl.rows[row_idx].cells[0].paragraphs[0].runs[0].bold = True
        set_cell(p2_tbl.rows[row_idx].cells[1], "", s1, font_size=8.8, space_before=4, space_after=4, line_spacing=1.18)
        set_cell(p2_tbl.rows[row_idx].cells[2], "", s2, font_size=8.8, space_before=4, space_after=4, line_spacing=1.18)
        set_cell(p2_tbl.rows[row_idx].cells[3], "", s3, font_size=8.8, space_before=4, space_after=4, line_spacing=1.18)

    add_p("[✓] Segment #1 Fear: The constant fear of losing walk-in patients and revenue simply because the lobby looks too chaotic.\n"
          "[✓] Segment #2 Frustration: The daily frustration of walking in blind and losing 1 to 2 valuable hours to stagnant queues.",
          bold_prefix="8. Circle the 2 that are the most intense for your customers:\n",
          space_before=8, space_after=7, line_spacing=1.22)

    add_p("[☑] Potential Segment #1: Independent Local Clinic & OPD Operators\n"
          "[☑] Potential Segment #2: Time-Constrained Students & Urban Patrons\n"
          "[☐] Potential Segment #3: Neighborhood Salons & Repair Shops",
          bold_prefix="9. Put a check next to the 2 you are most excited about:\n",
          space_before=5, space_after=7, line_spacing=1.22)

    add_p("• \"My ideal customers are independent neighborhood clinic and OPD doctors. I want to help them stop waiting-room chaos and keep their walk-in patients by letting them show live queue updates and issue simple virtual tokens.\"\n"
          "• \"My ideal customers are busy college students and working professionals. I want to help them stop wasting 1 to 2 hours in waiting rooms by letting them check line length and join the queue from their phone before leaving home.\"\n"
          "• \"My ideal customers are local walk-in business owners. I want to help them spread out customer rushes and prevent walk-outs without forcing anyone to use rigid appointment apps or buy expensive hardware.\"",
          bold_prefix="10. Positioning Statements (Pick 3):\n",
          space_before=5, space_after=4, line_spacing=1.22)

    doc.add_page_break()

    # ==========================================================
    # PAGE 3: THE PROBLEM STATEMENT CANVAS
    # ==========================================================
    add_sec_title("PAGE 3: THE PROBLEM STATEMENT CANVAS")

    meta_canvas = doc.add_table(rows=1, cols=3)
    meta_canvas.style = 'Table Grid'
    meta_canvas.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(meta_canvas.rows[0].cells[0], "STARTUP GROUPNAME: ", "CityPulse", font_size=10, space_before=4, space_after=4)
    set_cell(meta_canvas.rows[0].cells[1], "DATE: ", "15th Oct 2026", font_size=10, space_before=4, space_after=4)
    set_cell(meta_canvas.rows[0].cells[2], "VERSION: ", "1.0", font_size=10, space_before=4, space_after=4)

    # 3x3 Canvas Grid
    canvas_tbl = doc.add_table(rows=3, cols=3)
    canvas_tbl.style = 'Table Grid'
    canvas_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Row 1: Context, Problem, Alternatives
    set_cell(canvas_tbl.rows[0].cells[0], "CONTEXT\n(When does it occur?)\n",
             "During daily peak hours—especially weekday evenings (5:30 PM to 8:30 PM) and weekend mornings (10:00 AM to 1:00 PM)—when unannounced customer demand surges at local neighborhood service counters.",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)
    set_cell(canvas_tbl.rows[0].cells[1], "PROBLEM\n(Root Cause)\n",
             "Zero communication between offline service counters and prospective customers. Businesses have no friction-free way to show live capacity, and customers cannot check the line without physically walking through the door.",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)
    set_cell(canvas_tbl.rows[0].cells[2], "ALTERNATIVES\n(What customers do now)\n",
             "1. Physically showing up and waiting on crowded benches.\n2. Calling shop phones (rarely answered during busy hours).\n3. Looking at Google Maps static charts.\n4. Giving up and walking away.",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)

    # Row 2: Customers, Emotional Impact, Alternative Shortcomings
    set_cell(canvas_tbl.rows[1].cells[0], "CUSTOMERS\n(Who has it most?)\n",
             "Independent clinic doctors, OPD receptionists, and salon operators (Business Side); college students and working professionals on tight daily routines (Consumer Side).",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)
    set_cell(canvas_tbl.rows[1].cells[1], "EMOTIONAL IMPACT\n(How they feel)\n",
             "Customers: Anxious about lost time, annoyed by false 'just 10 minutes' estimates, physically drained.\n\nMerchants: Stressed by lobby noise, overwhelmed by inquiries, and frustrated watching walk-in sales leave.",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)
    set_cell(canvas_tbl.rows[1].cells[2], "ALTERNATIVE SHORTCOMINGS\n(Disadvantages)\n",
             "• In-person waiting wastes hours and creates health risks in clinics.\n• Phone calls interrupt ongoing service.\n• Google Maps only shows outdated averages.\n• Booking apps (e.g. Practo) enforce rigid slots and charge high fees.",
             font_size=9, space_before=5, space_after=5, line_spacing=1.2)

    # Row 3: Quantifiable Impact (Merged across 3 columns)
    r3 = canvas_tbl.rows[2].cells[0]
    r3.merge(canvas_tbl.rows[2].cells[1]).merge(canvas_tbl.rows[2].cells[2])
    set_cell(r3, "QUANTIFIABLE IMPACT (Measurable Impact with Units):\n",
             "• 45 to 75 minutes of unproductive idle waiting time per customer per visit.\n"
             "• 15% to 25% daily walk-in customer revenue lost by shops during peak hours due to overcrowded rooms.\n"
             "• 15 to 20 interrupted phone calls handled by the owner or receptionist in a single 3-hour evening shift.",
             font_size=9.2, space_before=5, space_after=5, line_spacing=1.22)

    # Detailed Market Validation Narrative Box (Ensures full vertical page coverage)
    add_p("• Core Problem Thesis: Offline service businesses remain completely disconnected from the digital schedules of their customers, leading to avoidable idle time and lost footfall.\n"
          "• Solution Validation: CityPulse bridges this operational asymmetry by allowing local shops to broadcast live waiting lines and issue lightweight virtual tokens with a single toggle.\n"
          "• Zero Friction Guarantee: Patrons do not need to download an app or create an account to view line status, and merchants require zero hardware investment.",
          bold_prefix="Problem Statement Canvas Analysis & Strategic Synthesis:\n",
          space_before=10, space_after=4, line_spacing=1.25)

    doc.add_page_break()

    # ==========================================================
    # PAGE 4: ACTIVITY — PILOT YOUR PURPOSE
    # ==========================================================
    add_sec_title("PAGE 4: ACTIVITY — PILOT YOUR PURPOSE")

    add_p("To stop people from wasting their day in crowded waiting rooms by creating a dead-simple, "
          "real-time link between local walk-in businesses and neighborhood customers: \"Know before you go.\"",
          bold_prefix="1. Venture Purpose & Ethos:\n",
          space_before=4, space_after=8, line_spacing=1.25)

    add_p("Every single day, people waste hours sitting on plastic chairs waiting for their turn, while independent "
          "business owners lose paying customers because their shops look overcrowded. Corporate chains have their own booking "
          "apps, but local neighborhood businesses are left behind. We believe getting a haircut or seeing a doctor should be "
          "predictable and hassle-free, without forcing anyone into rigid time slots or complex setups.",
          bold_prefix="2. Core Problem Conviction:\n",
          space_before=4, space_after=8, line_spacing=1.25)

    add_p("• We Live This Problem: As hostel students without personal transport, surprise 90-minute waits at local clinics or barbers throw off our academic schedules, lab timings, and hostel curfews.\n"
          "• Skills to Build It: As CSE students, we have the web-development background (Python, Flask, SQLite, Bootstrap) to build a fast, lightweight platform that runs in any mobile browser without forcing app downloads or merchant hardware costs.\n"
          "• Proximity to Market: Our campus is right beside busy student and residential hubs with dozens of clinics, salons, and repair counters within 3 km, enabling immediate real-world testing and daily feedback.",
          bold_prefix="3. Founder-Problem Fit (Why Our Team?):\n",
          space_before=4, space_after=8, line_spacing=1.25)

    add_p("• Keep It Dead Simple: A shop owner must update their line in five seconds with one click. A customer must see the wait time without downloading an app or making an account.\n"
          "• Realistic Times, Not Guesses: Wait-time predictions must dynamically adjust using moving averages based on real service times.\n"
          "• Support Local Businesses: Help neighborhood businesses keep their walk-in customers instead of losing them to large corporate chains.",
          bold_prefix="4. Guiding Operational Principles:\n",
          space_before=4, space_after=8, line_spacing=1.25)

    # 30-Day Pilot Execution Plan Table
    p4_tbl = doc.add_table(rows=4, cols=2)
    p4_tbl.style = 'Table Grid'
    p4_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    pilot_data = [
        ("Target Establishments", "3 local partner businesses: 1 neighborhood clinic/OPD, 1 busy barbershop/salon, and 1 local campus repair counter."),
        ("Wait Time Reduction", "Cut down physical lobby waiting time from 50 minutes to under 15 minutes for users using the virtual queue."),
        ("User Adoption", "Log at least 100 completed virtual queue check-ins over a 2-week active trial period."),
        ("Revenue Retention", "Reduce peak-hour customer walk-away rates for pilot shops from ~20% down to under 5%.")
    ]
    for i, (metric, target) in enumerate(pilot_data):
        set_cell(p4_tbl.rows[i].cells[0], metric, "", font_size=9.5, space_before=4, space_after=4)
        p4_tbl.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        set_cell(p4_tbl.rows[i].cells[1], "", target, font_size=9.2, space_before=4, space_after=4, line_spacing=1.2)

    # Save Document
    output_filename = "CityPulse_Entrepreneurship_Assignment.docx"
    doc.save(output_filename)
    print(f"Document successfully created and formatted: {output_filename}")

if __name__ == '__main__':
    create_full_page_assignment()